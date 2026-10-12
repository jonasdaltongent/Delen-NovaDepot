#!/usr/bin/env python3
"""
maak_werkdocumenten.py
Genereert het materiaal voor les 08 "Delen met de juiste rechten" (NovaDepot, ORLO).

  WP3_Delen-met-de-juiste-rechten.docx   het werkdocument dat de leerling INLEVERT
  Leveringsplanning.docx                 het oefendocument van de duo-test

Beide gaan als "Een kopie maken voor elke leerling" in de opdracht (soort "kopie" in classroom.json): elke
leerling krijgt een eigen Google-document en deelt de Leveringsplanning met de buur. Geen zip in deze les.

Gebruik:  python3 maak_werkdocumenten.py
Vereist:  python-docx

LET OP bij aanpassen:
 - De Leveringsplanning bestaat uit VIER LOSSE REGELS (gewone alinea's, geen tabel): in stap 4 wist de buur de
   regel van 10.30 uur, en dat moet kunnen met selecteren en Delete.
 - De leveringen komen uit les 05 en 06: Bakkerij Korstjes (7.30 uur, poort 3), Papierhandel Vellekens (8.15 uur),
   Koffiebranderij De Bonenbaas (10.30 uur) en Speelgoed Tolletje (13.00 uur, 3 pallets, poort 3), op dinsdag
   13 oktober 2026. De poorten en pallets van Vellekens en De Bonenbaas zijn voor deze les verzonnen.
 - De personages zijn die van de hele reeks: Tom Peeters (magazijn), Sara Claes (preventiedienst), Lotte Maes
   (personeelsdienst). Alle namen en bedrijven zijn fictief. Geen e-mailadressen, geen namen van leerlingen.
 - Rij 1 van deel 1 en deel 3 staat als voorbeeld ingevuld.
"""
import os

import huisstijl as hs   # de huisstijl van het werkdocument (uit _sjablonen/lespakket/werkdocument/)

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
UIT = os.path.join(HERE, "WP3_Delen-met-de-juiste-rechten.docx")
PLANNING = os.path.join(HERE, "Leveringsplanning.docx")

NAVY = RGBColor(0x2A, 0x39, 0x73)
GREY = RGBColor(0x4D, 0x55, 0x73)

# De vier leveringen van dinsdag 13 oktober 2026: uur, leverancier, poort, pallets
LEVERINGEN = [
    ("7.30 uur", "Bakkerij Korstjes", "poort 3", "2 pallets"),
    ("8.15 uur", "Papierhandel Vellekens", "poort 1", "1 pallet"),
    ("10.30 uur", "Koffiebranderij De Bonenbaas", "poort 2", "1 pallet"),
    ("13.00 uur", "Speelgoed Tolletje", "poort 3", "3 pallets"),
]


# ---------- hulpfuncties (zoals in les 06 en 07) ----------
def lettertype(run, naam="Arial"):
    run.font.name = naam
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), naam)


def basis_document(titel):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    rpr = st.element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), "Arial")
    taal = OxmlElement("w:lang")
    taal.set(qn("w:val"), "nl-BE")
    rpr.append(taal)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.1
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2)
    doc.core_properties.title = titel
    doc.core_properties.author = "Toegepaste Informatica"
    return doc


def tekst(doc, s, vet=False, klein=False, cursief=False, na=6, grootte=11, kleur=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(na)
    r = p.add_run(s)
    lettertype(r)
    r.bold = vet
    r.italic = cursief
    r.font.size = Pt(9.5 if klein else grootte)
    if klein:
        r.font.color.rgb = GREY
    if kleur is not None:
        r.font.color.rgb = kleur
    return p


# ---------- 1. Het werkdocument ----------
def werkdocument():
    """In de huisstijl van 12-10-2026 (huisstijl.py): alleen wat de leerling invult. Hoe hij werkt, staat op
    de lespagina; de zelfcontrole is de checklist van de lespagina."""
    doc = hs.nieuw_document()
    doc.core_properties.title = "WP3 Delen met de juiste rechten - werkdocument"
    hs.kopblok(doc, "DELEN MET DE JUISTE RECHTEN", punten=20)

    hs.deel(doc, 1, "Wat kan je met elke rol?")
    hs.tabel(doc, ["Wat kan je?", "Als Kijker", "Als Reageerder", "Als Bewerker"], [
        ("De planning lezen", "ja", "ja", "ja"),
        ("Een reactie schrijven", "", "", ""),
        ("Een regel zelf veranderen (een voorstel telt niet)", "", "", ""),
    ], [6.0, 3.6, 3.7, 3.7], invul=(1, 2, 3), rijhoogte=1.0, voorbeeld_rij=1)

    hs.deel(doc, 2, "Een regel terug")
    hs.vraag(doc, "Vraag 1. Welke regel wiste je buur in jouw planning?")
    hs.antwoordvak(doc, 1.2)
    hs.vraag(doc, "Vraag 2. Hoe kreeg je de regel terug? Schrijf het in je eigen woorden.")
    hs.antwoordvak(doc, 2.4)

    hs.deel(doc, 3, "Wie krijgt welke rol?")
    hs.tabel(doc, ["Wie?", "Wat moet die persoon kunnen?", "Welke rol?"], [
        ("Sara Claes (preventiedienst)", "de planning lezen, maar niets veranderen", "Kijker"),
        ("Lotte Maes (personeelsdienst)", "de planning lezen en een opmerking schrijven", ""),
        ("Een collega in het magazijn", "de planning lezen en aanpassen", ""),
        ("Je leraar, voor je map NovaDepot", "je map bekijken", ""),
    ], [4.8, 8.2, 4.0], invul=(2,), rijhoogte=1.0, voorbeeld_rij=1)
    hs.vraag(doc, "Waarom kies je bij Algemene toegang Beperkt, en niet Iedereen met de link?")
    hs.antwoordvak(doc, 1.8)

    hs.deel(doc, None, "Extra — niet verplicht")
    hs.vraag(doc, "a) Geef de herstelde versie van je planning een naam, bijvoorbeeld Planning in orde. "
                  "Welke naam gaf je?")
    hs.antwoordvak(doc, 1.2)
    hs.vraag(doc, "b) Beantwoord de reactie van je buur in je planning. Wat antwoordde je?")
    hs.antwoordvak(doc, 1.2)

    doc.save(UIT)


# ---------- 2. De Leveringsplanning ----------
def leveringsplanning():
    doc = basis_document("Leveringsplanning")
    tekst(doc, "LEVERINGSPLANNING", vet=True, grootte=24, kleur=NAVY, na=2)
    tekst(doc, "Dinsdag 13 oktober 2026", vet=True, grootte=16, na=0)
    tekst(doc, "NovaDepot  ·  magazijn  ·  opgemaakt door Tom Peeters  ·  fictief bedrijf", klein=True, na=16)
    for uur, wie, poort, pallets in LEVERINGEN:
        # elke levering is één gewone regel: selecteren en Delete volstaat om ze te wissen
        tekst(doc, f"{uur}  ·  {wie}  ·  {poort}  ·  {pallets}", grootte=13, na=10)
    tekst(doc, "")
    tekst(doc, "Vragen of opmerkingen? Schrijf ze als reactie in dit document.", klein=True)
    doc.save(PLANNING)


if __name__ == "__main__":
    werkdocument()
    leveringsplanning()
    print("Klaar.")
    print("  werkdocument om in te leveren : " + os.path.basename(UIT))
    print("  oefendocument (duo-test)      : " + os.path.basename(PLANNING))
